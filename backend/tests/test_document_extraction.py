from pathlib import Path

from app.services.document_extraction import (
    DocumentExtractionService,
)


def test_extract_txt(tmp_path):
    """Verify text extraction from TXT files."""

    file_path = tmp_path / "sample.txt"

    file_path.write_text(
        "Employee Handbook\nCompany policies",
        encoding="utf-8",
    )

    service = DocumentExtractionService()

    result = service.extract(str(file_path))

    assert "Employee Handbook" in result
    assert "Company policies" in result


def test_extract_pdf(tmp_path):
    """Verify PDF extraction."""

    from pypdf import PdfWriter

    file_path = tmp_path / "sample.pdf"

    writer = PdfWriter()

    writer.add_blank_page(
        width=612,
        height=792,
    )

    with open(file_path, "wb") as file:
        writer.write(file)

    service = DocumentExtractionService()

    result = service.extract(str(file_path))

    assert isinstance(result, str)


def test_extract_docx(tmp_path):
    """Verify DOCX extraction."""

    from docx import Document

    file_path = tmp_path / "sample.docx"

    document = Document()

    document.add_paragraph(
        "Employee Handbook"
    )

    document.add_paragraph(
        "Company policies"
    )

    document.save(file_path)

    service = DocumentExtractionService()

    result = service.extract(str(file_path))

    assert "Employee Handbook" in result
    assert "Company policies" in result


def test_unsupported_file_type(tmp_path):
    """Verify unsupported files are rejected."""

    file_path = tmp_path / "sample.xyz"

    file_path.write_text(
        "Unsupported file",
        encoding="utf-8",
    )

    service = DocumentExtractionService()

    try:
        service.extract(str(file_path))
        assert False
    except ValueError as exc:
        assert "Unsupported document type" in str(exc)
